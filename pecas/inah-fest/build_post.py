"""Inah Fest — post 1080x1350 seguindo o Manual de Identidade Visual do Inah.

Gera:
  out/fundo-malha-fogueira.png   (Malha Fogueira, definição do manual 4.2)
  out/fitas-fogueira.png         (fita-título + fita fina, degradê Fogueira 4.1/4.3/4.4)
  out/inah-fest.pptx             (texto editável com as fontes oficiais, para importar no Canva)
  out/preview.png                (prévia local renderizada com as fontes oficiais)
"""
import math, os, sys
import numpy as np
from PIL import Image, ImageDraw, ImageFont, ImageFilter
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

KIT = "/home/claude/brand-kit"
IMG = f"{KIT}/manual-identidade-visual/img"
ICO = f"{KIT}/manual-identidade-visual/icones/png"
FONTS = os.path.join(os.path.dirname(__file__), "fonts")
OUT = os.path.join(os.path.dirname(__file__), "out")
os.makedirs(OUT, exist_ok=True)

W, H = 1080, 1350
C = {
    "noite": "#0E1433", "inah": "#1D2756", "laranja": "#FF7A3D", "lanterna": "#FFD447",
    "verde": "#70EF8E", "branco": "#FFFFFF", "ceu": "#2F6FD8",
}
def rgb(h): h = h.lstrip("#"); return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))

# ------------------------------------------------------------------ malha
def malha_fogueira(w, h):
    """CSS do manual (build.py, MESH['fogueira']):
    radial-gradient(90% 55% at 50% 108%, Lanterna 1 0%, transp 55%),
    radial-gradient(110% 70% at 50% 110%, Laranja .95 0%, transp 70%),
    linear-gradient(180deg, #0E1433 0%, #1D2756 70%)"""
    y, x = np.mgrid[0:h, 0:w].astype(np.float64)
    t = np.clip(y / (0.70 * h), 0, 1)[..., None]
    base = (1 - t) * np.array(rgb(C["noite"])) + t * np.array(rgb(C["inah"]))
    def layer(rx, ry, cx, cy, col, a, stop):
        d = np.sqrt(((x - cx * w) / (rx * w)) ** 2 + ((y - cy * h) / (ry * h)) ** 2)
        al = a * np.clip(1 - d / stop, 0, 1)
        return np.array(rgb(col)), al[..., None]
    out = base
    # camadas de baixo para cima (a primeira do CSS fica por cima)
    for rx, ry, cx, cy, col, a, stop in [(1.10, .70, .5, 1.10, C["laranja"], .95, .70),
                                          (0.90, .55, .5, 1.08, C["lanterna"], 1.0, .55)]:
        c, al = layer(rx, ry, cx, cy, col, a, stop)
        out = out * (1 - al) + c * al
    return Image.fromarray(np.clip(out + 0.5, 0, 255).astype(np.uint8), "RGB")

# ------------------------------------------------------------------ fitas
SS = 3  # supersampling
def band_points(p0, p1, p2, p3, half, n=200):
    """Contorno de uma faixa com largura constante ao longo de uma Bézier cúbica."""
    pts, nor = [], []
    for i in range(n + 1):
        t = i / n
        b = [(1 - t) ** 3 * p0[k] + 3 * (1 - t) ** 2 * t * p1[k] + 3 * (1 - t) * t ** 2 * p2[k] + t ** 3 * p3[k] for k in (0, 1)]
        d = [3 * (1 - t) ** 2 * (p1[k] - p0[k]) + 6 * (1 - t) * t * (p2[k] - p1[k]) + 3 * t ** 2 * (p3[k] - p2[k]) for k in (0, 1)]
        L = math.hypot(*d); nx, ny = -d[1] / L, d[0] / L
        pts.append(b); nor.append((nx, ny))
    top = [(p[0] + nx * half, p[1] + ny * half) for p, (nx, ny) in zip(pts, nor)]
    bot = [(p[0] - nx * half, p[1] - ny * half) for p, (nx, ny) in zip(pts, nor)]
    return top + bot[::-1]

def grad_fogueira(w, h):
    """Degradê Fogueira 135°: Amarelo Lanterna 0% -> Laranja Fogueira 100%."""
    y, x = np.mgrid[0:h, 0:w].astype(np.float64)
    t = np.clip((x + y) / (w + h), 0, 1)[..., None]
    g = (1 - t) * np.array(rgb(C["lanterna"])) + t * np.array(rgb(C["laranja"]))
    return Image.fromarray(np.clip(g + .5, 0, 255).astype(np.uint8), "RGB")

# fita-título: curva suave, sangra pelas duas laterais, ~118 px de altura (≈ 2x a altura das letras)
FITA = dict(p0=(-80, 662), p1=(300, 640), p2=(760, 560), p3=(1160, 556), half=60)
# fita fina de apoio (trilha paralela), largura diferente, sem texto
FITA2 = dict(p0=(-80, 748), p1=(330, 726), p2=(770, 646), p3=(1160, 640), half=11)

def fitas(w, h):
    ws, hs = w * SS, h * SS
    mask = Image.new("L", (ws, hs), 0); d = ImageDraw.Draw(mask)
    for f in (FITA, FITA2):
        sc = lambda p: (p[0] * SS, p[1] * SS)
        d.polygon(band_points(sc(f["p0"]), sc(f["p1"]), sc(f["p2"]), sc(f["p3"]), f["half"] * SS), fill=255)
    mask = mask.resize((w, h), Image.LANCZOS)
    g = grad_fogueira(w, h).convert("RGBA"); g.putalpha(mask)
    return g

def fita_angle():
    p0, p3 = FITA["p0"], FITA["p3"]
    return math.degrees(math.atan2(p3[1] - p0[1], p3[0] - p0[0]))

# ------------------------------------------------------------------ layout
M = 72  # margem
ang = fita_angle()
def curve_y_at(f, xq):
    best=None
    for i in range(2001):
        t=i/2000; p=[(1-t)**3*f["p0"][k]+3*(1-t)**2*t*f["p1"][k]+3*(1-t)*t**2*f["p2"][k]+t**3*f["p3"][k] for k in (0,1)]
        if best is None or abs(p[0]-xq)<abs(best[0]-xq): best=p
    return best[1]
FITA_CY = curve_y_at(FITA, 540)
FITA_TXT_ROT = round(ang, 1)

# Elementos: imagens, formas e textos. Coordenadas em px no quadro 1080x1350.
IMAGES = [
    # flor sangrada: original, verde 100%, >= 1/3 da largura, sai pelo topo e pela direita (6.1)
    dict(src=f"{IMG}/flor.png", x=628, y=-128, w=560, h=560, name="Flor de lis"),
    # ícones oficiais (versão padrão, mancha Verde Inah), mesmo tamanho (7.6)
    dict(src=f"{ICO}/calendario.png", x=102, y=1086, w=76, h=76, name="Ícone calendário"),
    dict(src=f"{ICO}/sede.png", x=712, y=1086, w=76, h=76, name="Ícone sede"),
]
# tarjeta horizontal na base (2.7): marca 300 x 66 (x = 66), margem x/2 = 33, sai pela base
LOGO_W = 300; LOGO_H = round(LOGO_W * 126 / 575)
TJ_M = LOGO_H / 2
TJ_W = LOGO_W + 2 * TJ_M
TJ_X = (W - TJ_W) / 2
TJ_Y = H - (LOGO_H + TJ_M)          # topo da tarjeta; o branco passa do corte
IMAGES.append(dict(src=f"{IMG}/horizontal-positivo.png", x=TJ_X + TJ_M, y=TJ_Y + TJ_M, w=LOGO_W, h=LOGO_H,
                   name="Marca horizontal positiva"))

SHAPES = [
    # placares
    dict(kind="rrect", x=M, y=820, w=452, h=232, r=24, fill=C["inah"], line=C["verde"], lw=4, name="Placar barraquinha"),
    dict(kind="rrect", x=W - M - 452, y=820, w=452, h=232, r=24, fill=C["inah"], line=C["verde"], lw=4, name="Placar arquibancada"),
    dict(kind="rrect", x=M + 28, y=796, w=300, h=50, r=12, fill=C["verde"], name="Etiqueta barraquinha"),
    dict(kind="rrect", x=W - M - 452 + 28, y=796, w=330, h=50, r=12, fill=C["verde"], name="Etiqueta arquibancada"),
    # ingresso (Azul Noite sobre o brilho da malha)
    dict(kind="rrect", x=M, y=1070, w=W - 2 * M, h=148, r=24, fill=C["noite"], name="Ingresso"),
    dict(kind="dash", x=692, y=1086, h=116, color=C["lanterna"], name="Picote"),
    # tarjeta branca da marca saindo pela base (cantos de baixo passam do corte)
    dict(kind="rrect", x=TJ_X, y=TJ_Y, w=TJ_W, h=LOGO_H + TJ_M + 40 + TJ_M, r=TJ_M, fill=C["branco"], name="Tarjeta"),
]

BALOO, BARLOW, UBU = "Baloo 2", "Barlow Condensed", "Ubuntu"
TEXTS = [
    dict(t="O INAH CONVIDA", x=M, y=88, w=520, size=40, font=BARLOW, weight="ExtraBold", color=C["verde"], track=8, lh=1.0),
    dict(t="Inah\nFest!", x=M - 6, y=126, w=600, size=196, font=BALOO, weight="ExtraBold", color=C["lanterna"], lh=0.84),
    # fita-título (4.4): Ubuntu Bold caixa alta, azul, espaçamento leve
    dict(t="LANCHE DE PORTA DE ESTÁDIO", x=0, y=FITA_CY - 31, w=W, size=50, font=UBU, weight="Bold", color=C["noite"],
         track=4, lh=1.0, align="center", rot=FITA_TXT_ROT),
    dict(t="NA BARRAQUINHA", x=M + 28, y=803, w=300, size=34, font=BARLOW, weight="ExtraBold", color=C["inah"],
         track=6, lh=1.0, align="center"),
    dict(t="NA ARQUIBANCADA", x=W - M - 452 + 28, y=803, w=330, size=34, font=BARLOW, weight="ExtraBold", color=C["inah"],
         track=6, lh=1.0, align="center"),
    dict(t="Cachorro-quente\nEspetinho\nBolos e muito mais!", x=M + 32, y=872, w=400, size=36, font=UBU, weight="Medium",
         color=C["branco"], lh=1.3),
    dict(t="Karaokê\nBingo\nBrincadeiras pra família", x=W - M - 452 + 32, y=872, w=400, size=36, font=UBU,
         weight="Medium", color=C["branco"], lh=1.3),
    # ingresso
    dict(t="SÁBADO · [DD/MM]", x=196, y=1088, w=480, size=64, font=BARLOW, weight="ExtraBold", color=C["lanterna"],
         track=2, lh=1.0),
    dict(t="A PARTIR DAS [XX]H", x=196, y=1158, w=480, size=40, font=BARLOW, weight="ExtraBold", color=C["branco"],
         track=4, lh=1.0),
    dict(t="NA SEDE DO INAH", x=804, y=1094, w=190, size=28, font=BARLOW, weight="ExtraBold", color=C["verde"],
         track=4, lh=1.0),
    dict(t="Rua Hungria, 8\nParque das Nações\nSanto André", x=804, y=1130, w=190, size=22, font=UBU,
         weight="Regular", color=C["branco"], lh=1.2),
]
# @ e site ficam nos cantos, onde o brilho da malha ainda é escuro: branco, alinhados à tarjeta
TEXTS.append(dict(t="@inah.147", x=M, y=TJ_Y + 30, w=TJ_X - M - 20, size=28, font=UBU, weight="Medium",
                  color=C["branco"], lh=1.0))
TEXTS.append(dict(t="gepim.com.br", x=TJ_X + TJ_W + 20, y=TJ_Y + 30, w=W - M - (TJ_X + TJ_W + 20), size=28,
                  font=UBU, weight="Medium", color=C["branco"], lh=1.0, align="right"))

# ------------------------------------------------------------------ preview (PIL)
FONTFILE = {
    (BARLOW, "ExtraBold"): "BarlowCondensed-ExtraBold.ttf", (BARLOW, "SemiBold"): "BarlowCondensed-SemiBold.ttf",
    (UBU, "Bold"): "Ubuntu-Bold.ttf", (UBU, "Medium"): "Ubuntu-Medium.ttf", (UBU, "Regular"): "Ubuntu-Regular.ttf",
    (UBU, "Light"): "Ubuntu-Light.ttf", (BALOO, "ExtraBold"): "Baloo2%5Bwght%5D.ttf",
}
def pil_font(t):
    f = ImageFont.truetype(os.path.join(FONTS, FONTFILE[(t["font"], t["weight"])]), t["size"])
    if t["font"] == BALOO:
        f.set_variation_by_name("ExtraBold")
    return f

def draw_text(canvas, t):
    f = pil_font(t)
    lines = t["t"].split("\n")
    track = t.get("track", 0) / 100 * t["size"]
    asc, desc = f.getmetrics()
    lh = t["size"] * t.get("lh", 1.2)
    widths = [sum(f.getlength(ch) + track for ch in ln) - track for ln in lines]
    bw = int(max(t["w"], max(widths) + 4)); bh = int(lh * (len(lines) - 1) + asc + desc + 4)
    layer = Image.new("RGBA", (bw, bh), (0, 0, 0, 0)); d = ImageDraw.Draw(layer)
    for i, ln in enumerate(lines):
        al = t.get("align", "left")
        x = 0 if al == "left" else (t["w"] - widths[i]) / 2 if al == "center" else t["w"] - widths[i]
        y = i * lh + (lh - (asc + desc)) / 2 if lh > (asc + desc) else i * lh
        for ch in ln:
            d.text((x, y), ch, font=f, fill=rgb(t["color"]))
            x += f.getlength(ch) + track
    if t.get("rot"):
        layer = layer.rotate(-t["rot"], expand=True, resample=Image.BICUBIC)
        cx, cy = t["x"] + t["w"] / 2, t["y"] + bh / 2
        canvas.alpha_composite(layer, (int(cx - layer.width / 2), int(cy - layer.height / 2)))
    else:
        canvas.alpha_composite(layer, (int(t["x"]), int(t["y"])))
    return widths, bh

def rrect_img(s):
    ws, hs = int(s["w"] * SS), int(s["h"] * SS)
    im = Image.new("RGBA", (ws, hs), (0, 0, 0, 0)); d = ImageDraw.Draw(im)
    lw = s.get("lw", 0) * SS
    if s.get("line"):
        d.rounded_rectangle([0, 0, ws - 1, hs - 1], s["r"] * SS, fill=rgb(s["line"]))
        d.rounded_rectangle([lw, lw, ws - 1 - lw, hs - 1 - lw], max(0, s["r"] * SS - lw), fill=rgb(s["fill"]))
    else:
        d.rounded_rectangle([0, 0, ws - 1, hs - 1], s["r"] * SS, fill=rgb(s["fill"]))
    return im.resize((int(s["w"]), int(s["h"])), Image.LANCZOS)

def render_preview(bg, fx):
    cv = bg.convert("RGBA")
    cv.alpha_composite(fx)
    for s in SHAPES:
        if s["kind"] == "rrect":
            im = rrect_img(s)
            cv.alpha_composite(im, (int(s["x"]), int(s["y"]))) if s["y"] + s["h"] <= H else \
                cv.alpha_composite(im.crop((0, 0, im.width, int(H - s["y"]))), (int(s["x"]), int(s["y"])))
        elif s["kind"] == "dash":
            d = ImageDraw.Draw(cv)
            y = s["y"]
            while y < s["y"] + s["h"]:
                d.rectangle([s["x"], y, s["x"] + 3, min(y + 11, s["y"] + s["h"])], fill=rgb(s["color"]))
                y += 20
    for im in IMAGES:
        src = Image.open(im["src"]).convert("RGBA").resize((int(im["w"]), int(im["h"])), Image.LANCZOS)
        cv.alpha_composite(src, (int(im["x"]), int(im["y"]))) if im["y"] >= 0 else \
            cv.alpha_composite(src.crop((0, int(-im["y"]), src.width, src.height)), (int(im["x"]), 0))
    for t in TEXTS:
        draw_text(cv, t)
    return cv.convert("RGB")

# ------------------------------------------------------------------ pptx
EMU = 9525
def px(v): return Emu(int(round(v * EMU)))

def build_pptx(bg_path, fx_path, path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(W), px(H)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    sl.shapes.add_picture(bg_path, 0, 0, px(W), px(H)).name = "Fundo Malha Fogueira"
    sl.shapes.add_picture(fx_path, 0, 0, px(W), px(H)).name = "Fitas Fogueira"
    for s in SHAPES:
        if s["kind"] == "rrect":
            sh = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(s["x"]), px(s["y"]), px(s["w"]), px(s["h"]))
            sh.adjustments[0] = min(0.5, s["r"] / min(s["w"], s["h"]))
            sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(*rgb(s["fill"]))
            if s.get("line"):
                sh.line.color.rgb = RGBColor(*rgb(s["line"])); sh.line.width = px(s["lw"])
            else:
                sh.line.fill.background()
            sh.shadow.inherit = False
            sh.name = s["name"]
        elif s["kind"] == "dash":
            ln = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(s["x"] + 2), px(s["y"]), px(s["x"] + 2), px(s["y"] + s["h"]))
            ln.line.color.rgb = RGBColor(*rgb(s["color"])); ln.line.width = px(4)
            ln.line._get_or_add_ln().append(ln.line._get_or_add_ln().makeelement(qn("a:prstDash"), {"val": "dash"}))
            ln.name = s["name"]
    for im in IMAGES:
        p = sl.shapes.add_picture(im["src"], px(im["x"]), px(im["y"]), px(im["w"]), px(im["h"]))
        p.name = im["name"]
    for t in TEXTS:
        _, bh = draw_text(Image.new("RGBA", (W, H)), t)
        tb = sl.shapes.add_textbox(px(t["x"]), px(t["y"]), px(t["w"]), px(bh))
        tb.name = t["t"].split("\n")[0][:30]
        tf = tb.text_frame; tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
        bodyPr = tf._txBody.find(qn("a:bodyPr"))
        for el in list(bodyPr):
            bodyPr.remove(el)
        bodyPr.append(bodyPr.makeelement(qn("a:noAutofit"), {}))
        for i, line in enumerate(t["t"].split("\n")):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[t.get("align", "left")]
            para.line_spacing = t.get("lh", 1.2)
            r = para.add_run(); r.text = line
            f = r.font
            # nome da família estática como no Google Fonts, para o Canva casar com a fonte da biblioteca
            fam = t["font"] if t["weight"] in ("Regular", "Bold") else f"{t['font']} {t['weight']}"
            f.name = fam; f.size = Pt(t["size"] * 0.75)
            f.bold = t["weight"] in ("Bold", "ExtraBold")
            f.color.rgb = RGBColor(*rgb(t["color"]))
            rPr = r._r.get_or_add_rPr()
            for tag in ("a:latin", "a:ea", "a:cs"):
                el = rPr.find(qn(tag))
                if el is None:
                    el = rPr.makeelement(qn(tag), {}); rPr.append(el)
                el.set("typeface", fam)
            if t.get("track"):
                rPr.set("spc", str(int(round(t["track"] / 100 * t["size"] * 0.75 * 100))))
        if t.get("rot"):
            tb.rotation = t["rot"] % 360
    prs.save(path)

if __name__ == "__main__":
    bg = malha_fogueira(W, H); bg.save(f"{OUT}/fundo-malha-fogueira.png")
    fx = fitas(W, H); fx.save(f"{OUT}/fitas-fogueira.png")
    prev = render_preview(bg, fx); prev.save(f"{OUT}/preview.png")
    build_pptx(f"{OUT}/fundo-malha-fogueira.png", f"{OUT}/fitas-fogueira.png", f"{OUT}/inah-fest.pptx")
    print("fita angle", round(ang, 2))
